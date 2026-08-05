import json
import uuid
from collections import Counter
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.db import transaction
from django.db.models import Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import generics, status
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.payments.models import CashflowTransaction, Wallet
from apps.services.models import GeneralSettings
from .models import CustomerGroup, ImportedCustomer, NotificationLog, SmsTemplate
from .serializers import (
    CustomerGroupSerializer,
    SimpleSmsSendSerializer,
    SmsCampaignSendSerializer,
    SmsTemplateSerializer,
)
from .services import (
    build_customer_summaries,
    ensure_default_sms_templates,
    extract_log_metadata,
    group_sms_batches,
    send_provider_sms,
    is_valid_iran_mobile,
    normalize_phone,
    sms_chars_per_segment,
    sms_cost_for_text,
    sms_price_per_segment,
    sms_segments_for_text,
)


CUSTOMER_IMPORT_PRICE = Decimal('500000')
CUSTOMER_IMPORT_HEADERS = [
    ('full_name', 'نام مشتری'),
    ('phone', 'شماره تلفن'),
    ('plate_number', 'پلاک'),
    ('car_model', 'مدل خودرو'),
    ('car_color', 'رنگ خودرو'),
    ('notes', 'توضیحات'),
]
CUSTOMER_IMPORT_HEADER_ALIASES = {
    'full_name': {'نام مشتری', 'نام', 'اسم', 'name', 'full name', 'full_name'},
    'phone': {'شماره تلفن', 'شماره موبایل', 'موبایل', 'تلفن', 'phone', 'mobile'},
    'plate_number': {'پلاک', 'شماره پلاک', 'plate', 'plate_number'},
    'car_model': {'مدل خودرو', 'خودرو', 'مدل', 'car model', 'car_model'},
    'car_color': {'رنگ خودرو', 'رنگ', 'car color', 'car_color'},
    'notes': {'توضیحات', 'یادداشت', 'notes', 'note'},
}


def _cell_text(value):
    return str(value or '').strip()


def _normalize_header(value):
    return _cell_text(value).replace('\u200c', ' ').replace('_', ' ').strip().lower()


def _normalize_import_phone(value):
    phone = normalize_phone(value)
    if len(phone) == 10 and phone.startswith('9'):
        return f'0{phone}'
    return phone


def _parse_customer_import_workbook(uploaded_file):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({'file': f'فایل اکسل قابل خواندن نیست: {exc}'})

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], [{'row': 1, 'message': 'فایل خالی است.'}]

    raw_headers = [_normalize_header(value) for value in rows[0]]
    index_by_key = {}
    for key, aliases in CUSTOMER_IMPORT_HEADER_ALIASES.items():
        normalized_aliases = {_normalize_header(alias) for alias in aliases}
        for index, header in enumerate(raw_headers):
            if header in normalized_aliases:
                index_by_key[key] = index
                break

    if 'phone' not in index_by_key:
        return [], [{'row': 1, 'message': 'ستون شماره تلفن در فایل پیدا نشد.'}]

    parsed = []
    errors = []
    seen_phones = set()
    for row_number, row in enumerate(rows[1:], start=2):
        values = {}
        for key, index in index_by_key.items():
            values[key] = _cell_text(row[index] if index < len(row) else '')
        if not any(values.values()):
            continue

        phone = _normalize_import_phone(values.get('phone'))
        if not is_valid_iran_mobile(phone):
            errors.append({'row': row_number, 'message': 'شماره تلفن معتبر نیست و باید با 09 شروع شود.'})
            continue
        if phone in seen_phones:
            errors.append({'row': row_number, 'message': 'شماره تلفن در همین فایل تکراری است.'})
            continue
        seen_phones.add(phone)
        parsed.append({
            'row': row_number,
            'full_name': values.get('full_name') or 'مشتری بدون نام',
            'phone': phone,
            'plate_number': values.get('plate_number', ''),
            'car_model': values.get('car_model', ''),
            'car_color': values.get('car_color', ''),
            'notes': values.get('notes', ''),
        })
    return parsed, errors


def _default_wallet_for_update(tenant):
    wallet = Wallet.objects.select_for_update().filter(
        tenant=tenant,
        wallet_type=Wallet.WalletType.BANK,
        is_active=True,
    ).order_by('id').first()
    if wallet:
        return wallet
    wallet = Wallet.objects.create(
        tenant=tenant,
        name='کیف پول اصلی',
        wallet_type=Wallet.WalletType.BANK,
        balance=0,
        is_active=True,
    )
    return Wallet.objects.select_for_update().get(id=wallet.id)


class SmsWalletMixin:
    def _get_sms_wallets_for_update(self, tenant):
        wallets = list(
            Wallet.objects.select_for_update()
            .filter(
                tenant=tenant,
                wallet_type=Wallet.WalletType.SMS,
                is_active=True,
            )
            .order_by('id')
        )
        if wallets:
            return wallets

        wallet = Wallet.objects.create(
            tenant=tenant,
            name='کیف پول پیامک',
            wallet_type=Wallet.WalletType.SMS,
            balance=0,
            is_active=True,
        )
        return [wallet]

    def _sms_wallet_balance(self, tenant):
        return (
            Wallet.objects.filter(
                tenant=tenant,
                wallet_type=Wallet.WalletType.SMS,
                is_active=True,
            ).aggregate(
                total=Coalesce(Sum('balance'), Value(Decimal('0')))
            )['total']
            or Decimal('0')
        )

    def _debit_sms_wallets(self, tenant, amount, *, description, reference_type, created_by=None):
        remaining = Decimal(str(amount or 0))
        wallets = self._get_sms_wallets_for_update(tenant)
        for wallet in wallets:
            if remaining <= 0:
                break
            current_balance = Decimal(str(wallet.balance or 0))
            if current_balance <= 0:
                continue
            debit = min(current_balance, remaining)
            wallet.balance = current_balance - debit
            wallet.save(update_fields=['balance', 'updated_at'])
            CashflowTransaction.objects.create(
                tenant=tenant,
                wallet=wallet,
                direction=CashflowTransaction.Direction.OUT,
                amount=debit,
                description=description,
                reference_type=reference_type,
                created_by=created_by if getattr(created_by, 'is_authenticated', False) else None,
            )
            remaining -= debit

        if remaining > 0:
            raise ValidationError({'detail': 'موجودی کیف پول پیامک کافی نیست.'})


class SmsProviderMixin:
    provider_url = '/ws/v1/sms/simple'

    def _general_settings(self):
        tenant = getattr(self.request.user, 'tenant', None)
        if not tenant:
            return None
        return GeneralSettings.objects.filter(tenant=tenant).first()

    def _provider_config(self):
        api_key = str(
            getattr(settings, 'IRANPAYAMAK_API_KEY', '')
            or ''
        ).strip()
        line_number = str(
            getattr(settings, 'IRANPAYAMAK_LINE_NUMBER', '')
            or ''
        ).strip()
        base_url = str(
            getattr(settings, 'IRANPAYAMAK_BASE_URL', 'https://api.iranpayamak.com')
            or 'https://api.iranpayamak.com'
        ).rstrip('/')
        if not api_key or not line_number:
            raise ValidationError({'detail': 'تنظیمات سرویس پیامک کامل نیست.'})
        return {
            'api_key': api_key,
            'line_number': line_number,
            'base_url': base_url,
        }

    def _provider_message_text(self, provider_data, fallback='ارسال پیامک توسط سرویس تایید نشد.'):
        if isinstance(provider_data, dict):
            message = provider_data.get('message') or provider_data.get('messages')
            if isinstance(message, list):
                return ' | '.join(str(item) for item in message if item)
            if isinstance(message, dict):
                return json.dumps(message, ensure_ascii=False)
            if message:
                try:
                    return bytes(str(message), 'utf-8').decode('unicode_escape')
                except Exception:
                    return str(message)
        if provider_data:
            return str(provider_data)
        return fallback

    def _send_provider_request(self, *, text, recipients):
        config = self._provider_config()
        return send_provider_sms(
            getattr(self.request.user, 'tenant', None),
            text,
            recipients,
            provider_config=config,
        )


class CustomerClubDashboardView(APIView, SmsWalletMixin):
    def get(self, request):
        tenant = getattr(request.user, 'tenant', None)
        ensure_default_sms_templates(tenant, request.user)
        customers = build_customer_summaries(tenant)
        templates = SmsTemplate.objects.filter(tenant=tenant, is_active=True).order_by('display_order', 'title', 'id')
        groups = CustomerGroup.objects.filter(tenant=tenant, is_active=True).order_by('-created_at', '-id')
        logs = NotificationLog.objects.filter(tenant=tenant, channel=NotificationLog.Channel.SMS).order_by('-created_at')[:100]
        status_counts = Counter(log.status for log in logs)

        return Response(
            {
                'summary': {
                    'sms_balance': self._sms_wallet_balance(tenant),
                    'sms_price_per_segment': float(sms_price_per_segment()),
                    'sms_chars_per_segment': sms_chars_per_segment(),
                    'status_counts': {
                        'success': status_counts.get(NotificationLog.Status.SENT, 0),
                        'pending': status_counts.get(NotificationLog.Status.PENDING, 0),
                        'failed': status_counts.get(NotificationLog.Status.FAILED, 0),
                    },
                },
                'customers': customers,
                'groups': CustomerGroupSerializer(groups, many=True).data,
                'templates': SmsTemplateSerializer(templates, many=True).data,
                'logs': [extract_log_metadata(log) for log in logs],
            }
        )


class CustomerImportTemplateView(APIView):
    def get(self, request):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'customers'
        sheet.append([label for _, label in CUSTOMER_IMPORT_HEADERS])
        sheet.append([
            'علی رضایی',
            '09123456789',
            '12 ب 345 67',
            'پژو 206',
            'سفید',
            'مشتری وفادار',
        ])
        sheet.append([
            'سارا احمدی',
            '09120000000',
            '',
            'تیبا',
            'مشکی',
            '',
        ])
        for index, width in enumerate([22, 18, 18, 18, 14, 32], start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="customer-import-template.xlsx"'
        return response


class CustomerImportPreviewView(APIView):
    def post(self, request):
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({'file': ['فایل اکسل را انتخاب کنید.']}, status=status.HTTP_400_BAD_REQUEST)

        rows, errors = _parse_customer_import_workbook(uploaded_file)
        return Response({
            'preview': rows[:5],
            'valid_count': len(rows),
            'error_count': len(errors),
            'errors': errors[:20],
            'price': CUSTOMER_IMPORT_PRICE,
        })


class CustomerImportConfirmView(APIView):
    def post(self, request):
        tenant = getattr(request.user, 'tenant', None)
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({'file': ['فایل اکسل را انتخاب کنید.']}, status=status.HTTP_400_BAD_REQUEST)

        rows, errors = _parse_customer_import_workbook(uploaded_file)
        if errors:
            return Response(
                {
                    'detail': 'فایل هنوز خطای قابل اصلاح دارد.',
                    'errors': errors[:20],
                    'valid_count': len(rows),
                    'error_count': len(errors),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not rows:
            return Response({'detail': 'هیچ مشتری معتبری در فایل پیدا نشد.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            wallet = _default_wallet_for_update(tenant)
            wallet_balance = Decimal(str(wallet.balance or 0))
            if wallet_balance < CUSTOMER_IMPORT_PRICE:
                return Response(
                    {
                        'detail': 'موجودی کیف پول برای وارد کردن مشتریان کافی نیست.',
                        'required_amount': CUSTOMER_IMPORT_PRICE,
                        'wallet_balance': wallet_balance,
                    },
                    status=status.HTTP_402_PAYMENT_REQUIRED,
                )

            wallet.balance = wallet_balance - CUSTOMER_IMPORT_PRICE
            wallet.save(update_fields=['balance', 'updated_at'])
            CashflowTransaction.objects.create(
                tenant=tenant,
                wallet=wallet,
                direction=CashflowTransaction.Direction.OUT,
                amount=CUSTOMER_IMPORT_PRICE,
                description='وارد کردن مشتریان باشگاه مشتریان از اکسل',
                reference_type='customer_import_excel',
                created_by=request.user if getattr(request.user, 'is_authenticated', False) else None,
            )

            created_count = 0
            updated_count = 0
            for item in rows:
                _, created = ImportedCustomer.objects.update_or_create(
                    tenant=tenant,
                    phone=item['phone'],
                    defaults={
                        'full_name': item['full_name'],
                        'plate_number': item.get('plate_number', ''),
                        'car_model': item.get('car_model', ''),
                        'car_color': item.get('car_color', ''),
                        'notes': item.get('notes', ''),
                        'source': 'excel',
                        'imported_by': request.user if getattr(request.user, 'is_authenticated', False) else None,
                    },
                )
                if created:
                    created_count += 1
                else:
                    updated_count += 1

        return Response({
            'created_count': created_count,
            'updated_count': updated_count,
            'imported_count': created_count + updated_count,
            'charged_amount': CUSTOMER_IMPORT_PRICE,
        }, status=status.HTTP_201_CREATED)


class CustomerGroupListCreateView(generics.ListCreateAPIView):
    serializer_class = CustomerGroupSerializer

    def get_queryset(self):
        return CustomerGroup.objects.filter(
            tenant=self.request.user.tenant,
            is_active=True,
        ).order_by('-created_at', '-id')

    def perform_create(self, serializer):
        serializer.save(
            tenant=self.request.user.tenant,
            created_by=self.request.user if getattr(self.request.user, 'is_authenticated', False) else None,
            updated_by=self.request.user if getattr(self.request.user, 'is_authenticated', False) else None,
        )


class CustomerGroupDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = CustomerGroupSerializer

    def get_queryset(self):
        return CustomerGroup.objects.filter(tenant=self.request.user.tenant, is_active=True)

    def perform_update(self, serializer):
        serializer.save(
            updated_by=self.request.user if getattr(self.request.user, 'is_authenticated', False) else None,
        )

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.updated_by = self.request.user if getattr(self.request.user, 'is_authenticated', False) else None
        instance.save(update_fields=['is_active', 'updated_by', 'updated_at'])


class SmsTemplateListCreateView(generics.ListCreateAPIView):
    serializer_class = SmsTemplateSerializer

    def get_queryset(self):
        tenant = self.request.user.tenant
        ensure_default_sms_templates(tenant, self.request.user)
        return SmsTemplate.objects.filter(tenant=tenant, is_active=True).order_by('display_order', 'title', 'id')

    def perform_create(self, serializer):
        serializer.save(
            tenant=self.request.user.tenant,
            created_by=self.request.user if getattr(self.request.user, 'is_authenticated', False) else None,
        )


class SmsTemplateDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = SmsTemplateSerializer

    def get_queryset(self):
        return SmsTemplate.objects.filter(tenant=self.request.user.tenant, is_active=True)

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.save(update_fields=['is_active', 'updated_at'])


class SmsCampaignSendView(APIView, SmsWalletMixin, SmsProviderMixin):
    def _create_log_entries(self, *, tenant, request_user, recipients, status_value, provider_message_id='', provider_response='', extra_payload=None):
        extra_payload = extra_payload or {}
        sent_at = timezone.now() if status_value == NotificationLog.Status.SENT else None
        logs = []
        for recipient in recipients:
            payload = {
                **extra_payload,
                'recipient_name': recipient.get('name', ''),
                'customer': {
                    'key': recipient.get('key', ''),
                    'name': recipient.get('name', ''),
                    'phone': recipient.get('phone', ''),
                    'carwash_name': recipient.get('carwash_name', ''),
                    'orders_count': recipient.get('orders_count', 0),
                    'total_spent': recipient.get('total_spent', 0),
                    'score': recipient.get('score', 0),
                    'primary_plate': recipient.get('primary_plate', ''),
                },
                'rendered_text': recipient.get('rendered_text', ''),
            }
            logs.append(
                NotificationLog(
                    tenant=tenant,
                    channel=NotificationLog.Channel.SMS,
                    recipient=recipient.get('phone', ''),
                    template_code=extra_payload.get('template_code', ''),
                    payload=payload,
                    status=status_value,
                    sent_at=sent_at,
                    provider_message_id=str(provider_message_id or ''),
                    provider_response=provider_response,
                    created_by=request_user if getattr(request_user, 'is_authenticated', False) else None,
                )
            )
        NotificationLog.objects.bulk_create(logs)

    def _send_validated_campaign(self, request, validated):
        tenant = getattr(request.user, 'tenant', None)
        recipients = validated['recipients']
        template_text = validated['template_text']
        sms_price = sms_price_per_segment()
        chars_per_segment = sms_chars_per_segment()
        campaign_id = str(uuid.uuid4())
        grouped_batches, _rendered_recipients = group_sms_batches(recipients, template_text, getattr(tenant, 'name', ''))

        estimated_total = sum(
            (sms_cost_for_text(rendered_text) * Decimal(len(batch_recipients)) for rendered_text, batch_recipients in grouped_batches.items()),
            Decimal('0'),
        )

        recipient_count = sum(len(batch_recipients) for batch_recipients in grouped_batches.values())
        try:
            from apps.auth.sample_tenant import assert_sample_sms_capacity, maybe_raise_sample_monthly_sms_alert

            assert_sample_sms_capacity(tenant, extra=recipient_count)
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        wallet_balance = self._sms_wallet_balance(tenant)
        if wallet_balance < estimated_total:
            return Response(
                {
                    'detail': 'موجودی کیف پول پیامک کافی نیست.',
                    'required_amount': estimated_total,
                    'wallet_balance': wallet_balance,
                },
                status=status.HTTP_402_PAYMENT_REQUIRED,
            )

        results = []
        success_count = 0
        failed_count = 0
        debited_amount = Decimal('0')

        for rendered_text, batch_recipients in grouped_batches.items():
            phone_numbers = [recipient['phone'] for recipient in batch_recipients]
            provider_result = self._send_provider_request(text=rendered_text, recipients=phone_numbers)
            unit_cost = sms_cost_for_text(rendered_text)
            segments = sms_segments_for_text(rendered_text)
            batch_cost = unit_cost * Decimal(len(batch_recipients))
            extra_payload = {
                'campaign_id': campaign_id,
                'target_label': validated.get('target_label', ''),
                'note': validated.get('note', ''),
                'template_code': validated.get('template_code', ''),
                'template_text': template_text,
                'segments': segments,
                'chars_per_segment': chars_per_segment,
                'price_per_sms': str(unit_cost),
                'price_per_segment': str(sms_price),
                'total_cost': str(batch_cost),
                'provider_request': provider_result['payload'],
            }

            if provider_result['ok']:
                with transaction.atomic():
                    self._debit_sms_wallets(
                        tenant,
                        batch_cost,
                        description=f'ارسال پیامک: {validated.get("target_label", "").strip() or "باشگاه مشتریان"}',
                        reference_type='sms_campaign_send',
                        created_by=request.user,
                    )
                debited_amount += batch_cost
                self._create_log_entries(
                    tenant=tenant,
                    request_user=request.user,
                    recipients=batch_recipients,
                    status_value=NotificationLog.Status.SENT,
                    provider_message_id=provider_result.get('provider_id', ''),
                    provider_response=provider_result['raw_body'],
                    extra_payload=extra_payload,
                )
                for recipient in batch_recipients:
                    success_count += 1
                    results.append(
                        {
                            'phone': recipient['phone'],
                            'recipient_name': recipient.get('name', ''),
                            'status': 'success',
                            'message': provider_result['message'],
                            'provider_message_id': provider_result.get('provider_id', ''),
                            'rendered_text': recipient.get('rendered_text', ''),
                            'cost': unit_cost,
                            'segments': segments,
                        }
                    )
            else:
                self._create_log_entries(
                    tenant=tenant,
                    request_user=request.user,
                    recipients=batch_recipients,
                    status_value=NotificationLog.Status.FAILED,
                    provider_response=provider_result['raw_body'],
                    extra_payload=extra_payload,
                )
                for recipient in batch_recipients:
                    failed_count += 1
                    results.append(
                        {
                            'phone': recipient['phone'],
                            'recipient_name': recipient.get('name', ''),
                            'status': 'failed',
                            'message': provider_result['message'],
                            'provider_message_id': '',
                            'rendered_text': recipient.get('rendered_text', ''),
                            'cost': unit_cost,
                            'segments': segments,
                        }
                    )

        http_status = status.HTTP_201_CREATED if success_count > 0 else status.HTTP_502_BAD_GATEWAY
        detail = 'پیامک با موفقیت در صف ارسال قرار گرفت.'
        if success_count and failed_count:
            detail = 'بخشی از پیامک‌ها ارسال شد و بخشی ناموفق بود.'
        elif failed_count and not success_count:
            detail = str(results[0].get('message') or 'هیچ پیامکی ارسال نشد.') if results else 'هیچ پیامکی ارسال نشد.'

        if success_count:
            from apps.auth.sample_tenant import maybe_raise_sample_monthly_sms_alert

            maybe_raise_sample_monthly_sms_alert()

        return Response(
            {
                'detail': detail,
                'campaign_id': campaign_id,
                'success_count': success_count,
                'failed_count': failed_count,
                'debited_amount': debited_amount,
                'results': results,
            },
            status=http_status,
        )

    def post(self, request):
        serializer = SmsCampaignSendSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        return self._send_validated_campaign(request, serializer.validated_data)


class SimpleSmsSendView(SmsCampaignSendView):
    def post(self, request):
        serializer = SimpleSmsSendSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        recipients = serializer.validated_data['recipients']
        payload = {
            'template_text': serializer.validated_data['text'],
            'target_label': serializer.validated_data.get('target_label', ''),
            'note': serializer.validated_data.get('note', ''),
            'recipients': [
                {
                    'phone': phone,
                    'name': '',
                    'carwash_name': getattr(getattr(request.user, 'tenant', None), 'name', ''),
                }
                for phone in recipients
            ],
        }
        campaign_serializer = SmsCampaignSendSerializer(data=payload)
        campaign_serializer.is_valid(raise_exception=True)
        return self._send_validated_campaign(request, campaign_serializer.validated_data)
